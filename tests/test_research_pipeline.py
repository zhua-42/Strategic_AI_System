"""Offline regression checks; fixtures are explicitly synthetic, not investment research."""
import io
import json
from pathlib import Path
import sys
import unittest
import zipfile

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from research_agent import proforma, validate_evidence, model_visual, quality, COMPANY_CHAPTERS
import research_documents as documents


def fixture():
    source={"id":"S1","title":"软件测试资料（非真实公司数据）","url":"https://example.invalid/fixture",
            "domain":"example.invalid","text":"测试公司2025年度营业收入为1,000百万元。所有数据均为软件测试。",
            "method":"离线测试夹具","retrieved_at":"2026-09-29T00:00:00Z"}
    facts=[{"subject":"测试公司","claim":"软件测试收入，不代表实际公司","metric":"revenue","value":1000,
            "annual":True,"is_target":True,"scope":"consolidated","period":"2025FY","currency":"CNY","unit":"百万元",
            "quote":source["text"],"source_id":"S1","topic":"收入"}]
    model=proforma(facts)
    paragraphs=[
      "本页为软件排版和公式验收样例，文中的测试公司与数值均不代表任何真实企业，不应作为投资资料使用。测试要验证研究段落、图表、来源说明和连续页码能否在同一份文档中正确排列。本段模拟券商研究报告的完整论证：先提出观察问题，再指出资料的适用范围，最后说明该问题如何传导到收入确认、成本结构和资金需求。实际生成时这些表述必须由已经打开的公告和财务报告支持，而不能从示例中复制。检索日期只是读取日期，与财务报表所覆盖的年度和季度属于不同概念，正式分析必须将二者分开列示。[S1]",
      "对于经营质量的分析，收入增长与利润增长需要结合回款和资本开支一起观察。短期订单变化可能受交付节奏影响，单一季度的高增速不能直接代表长期需求。比较不同公司时还需要统一币种、会计期间与统计范围，解释毛利和营业利润的差异。如果某项数据只出现在搜索结果摘要而未读取原文，就不能用于历史趋势图。当前测试只检查这些文字是否完整显示，不核实任何行业结论。缺少披露信息的维度应在报告中写清楚证据边界，再提出可以通过后续公告验证的研究假设，不能用无来源数字补齐。[S1]",
      "下方图表所对应的研究判断需要同时讨论正向驱动和反面条件。假设收入上升，但应收账款增长更快，经营现金流可能并未同步改善；假设投入扩大，也需要观察产能利用率和新增业务的现金回收周期。这些是模型机制而不是对测试公司经营状况的判断。预测列应明确标注字母E，并在图表下方列出资料编号以及计算依据。读者可以从证据工作簿查找原文引文，再修改假设工作表中的变量，查看结果是否符合财务报表的勾稽关系。任何缺口都应该保留为可追踪的问题。[S1]",
      "验证工作还包括权益与现金的两条桥梁：期末权益由期初权益、净利润、股东交易及分红构成，期末现金则由期初现金和经营、投资、筹资活动共同决定。现金不足时模型显式列出融资需求，不能悄悄调整资产使报表看起来平衡。本样例的期初结构和预测假设均已标识为测试情景。正式研究需要结合真实公告逐项替换，并检查金融类企业是否需要专用的监管资本和估值模型。报告的最终结论应当与事实覆盖范围相匹配，结构合格只说明文件满足排版和内容要求，并不意味着投资判断得到认证。[S1]"
    ]
    chapters=[]
    for i,title in enumerate(COMPANY_CHAPTERS):
        visual=model_visual(i,model)
        if not visual:
            visual={"kind":"table","title":title+"测试比较表", "headers":["测试维度","说明","资料"],
                "rows":[[title,"用于核对不同章节的正文与表格对应关系；不代表真实投研结论。","S1"],
                        ["期间","2025FY为测试期间，来源示例域名不应当被访问。","S1"],
                        ["局限","测试夹具仅验证文件格式、楷体嵌入和导出过程。","S1"],
                        ["反证条件","实际研究应当核对原文并披露相反的证据。","S1"]],
                "source":"来源：S1 软件测试夹具，非真实公司资料。"}
        if i in (9,10):
            visual={"kind":"bar" if i==9 else "line","title":"测试数列（非真实数据）", "labels":["2023测试","2024测试","2025测试"],
                    "values":[70,90,100] if i==9 else [-10,3,8],"unit":"测试单位","source":"来源：S1 合成测试数据，不作投资用途。"}
        chapters.append({"title":title,"paragraphs":paragraphs,"takeaway":"软件验收样例：验证图文交错、字体、来源说明和模型勾稽。","visual":visual})
    package={"target":"软件验收样例（非真实研报）","kind":"公司","report_type":"公司深度专题","period":"软件测试",
             "created_at":"2026-09-29T00:00:00Z","sources":[source],"facts":facts,"model":model,"chapters":chapters,"audit":[],"rejected":[]}
    package["checks"]=quality(package)
    package["status"]="软件测试文件，不构成真实投研报告"
    return package


class PipelineTests(unittest.TestCase):
    def test_four_statement_balances_under_stress(self):
        for assumptions in ({},{"growth":.25,"capex_ratio":.8},{"gross_margin":.05,"opex_ratio":.8},{"growth":-.4}):
            model=proforma([],overrides=assumptions)
            for row in model["rows"]:
                self.assertAlmostEqual(row["balance_check"],0,places=8)
                self.assertAlmostEqual(row["cash_check"],0,places=8)
                self.assertAlmostEqual(row["equity_check"],0,places=8)
                self.assertGreaterEqual(row["cash"],-1e-8)
        self.assertTrue(proforma([])["normalized"])

    def test_reject_fabricated_quote_and_number(self):
        p=fixture(); source=p["sources"][0]
        raw={"quote":source["text"],"source_id":"S1","claim":"test","raw_number":"999999"}
        facts,errors=validate_evidence({"facts":[raw,dict(raw,quote="This fabricated statement is absent")]},[source],"测试公司")
        self.assertEqual(len(facts),1)
        self.assertNotIn("value",facts[0])
        self.assertEqual(len(errors),1)

    def test_models_reject_invalid_discount_rate(self):
        with self.assertRaises(ValueError): proforma([],overrides={"wacc":.02,"terminal_growth":.03})

    def test_regional_revenue_cannot_be_company_base(self):
        source={"id":"S1","text":"比亚迪2025年境外营业收入为310,740,988千元。"}
        raw={"source_id":"S1","quote":source["text"],"claim":source["text"],"metric":"revenue","raw_number":"310,740,988",
             "scope":"consolidated","is_target":True,"annual":True,"period":"2025FY","unit":"千元","currency":"CNY"}
        facts,_=validate_evidence({"facts":[raw]},[source],"比亚迪")
        self.assertEqual(facts[0]["scope"],"segment")
        self.assertTrue(proforma(facts)["normalized"])

    def test_workbook_formulas_and_fonts(self):
        from openpyxl import load_workbook
        raw=documents.evidence_bytes(fixture())
        wb=load_workbook(io.BytesIO(raw),data_only=False)
        self.assertIn("Equity",wb.sheetnames)
        self.assertTrue(wb["Pro Forma"]["C2"].value.startswith("="))
        self.assertIn("NPV",wb["DCF Sensitivity"]["B2"].value)
        self.assertEqual(wb["Assumptions"]["B2"].font.name,"LXGW WenKai")
        vals=load_workbook(io.BytesIO(raw),data_only=True)
        self.assertAlmostEqual(vals["Pro Forma"]["G29"].value,0,places=6)
        self.assertAlmostEqual(vals["Equity"]["F8"].value,0,places=6)
        Path("tmp/pdfs/research-test.xlsx").write_bytes(raw)

    def test_pdf_word_ppt_outputs(self):
        from pypdf import PdfReader
        from pptx import Presentation
        package=fixture()
        raw=documents.pdf_bytes(package)
        Path("tmp/pdfs/research-test.pdf").write_bytes(raw)
        reader=PdfReader(io.BytesIO(raw))
        self.assertGreaterEqual(len(reader.pages),26)
        combined="".join(page.extract_text() for page in reader.pages)
        for chapter in package["chapters"]: self.assertIn(chapter["title"],combined)
        self.assertIn("来源：",combined)
        self.assertNotIn("待补充",combined)
        fonts=[]
        for page in reader.pages:
            for ref in page["/Resources"]["/Font"].values():
                obj=ref.get_object()
                if "/FontDescriptor" in obj: fonts.append(obj["/FontDescriptor"].get_object())
        self.assertTrue(any("/FontFile2" in obj for obj in fonts))
        docx=documents.docx_bytes(package)
        Path("tmp/pdfs/research-test.docx").write_bytes(docx)
        with zipfile.ZipFile(io.BytesIO(docx)) as archive:
            self.assertIn(b"LXGW WenKai",archive.read("word/styles.xml"))
        pptx=documents.pptx_bytes(package,raw)
        Path("tmp/pdfs/research-test.pptx").write_bytes(pptx)
        self.assertEqual(len(Presentation(io.BytesIO(pptx)).slides),len(reader.pages))
        print("Verified PDF/PPT page count:",len(reader.pages))


if __name__ == "__main__":
    Path("tmp/pdfs").mkdir(parents=True,exist_ok=True)
    unittest.main(verbosity=2)
