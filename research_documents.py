"""Kai-type, interleaved research documents shared by web, PDF, Word and PPT."""
import io
import json
from pathlib import Path
from html import escape

FONT_PATH = Path(__file__).parent / "assets/fonts/LXGWWenKai-Regular.ttf"
FONT_FAMILY = "LXGW WenKai"


def font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    if "ResearchKai" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("ResearchKai", str(FONT_PATH)))
    return "ResearchKai"


def chart_drawing(visual, width=480, height=185):
    from reportlab.graphics.shapes import Drawing, Rect, Line, String, Circle
    from reportlab.lib.colors import HexColor
    if visual["kind"] in ("flow","timeline"):
        rows=visual["rows"][:6]
        height=100*((len(rows)+1)//2)+10
        drawing=Drawing(width,height)
        for i,row in enumerate(rows):
            x=8+(i%2)*width/2
            y=height-100*(i//2+1)
            boxw=width/2-22
            drawing.add(Rect(x,y,boxw,88,fillColor=HexColor("#F0F4F8"),strokeColor=HexColor("#285F95"),strokeWidth=.6))
            drawing.add(String(x+8,y+73,f"{i+1}. "+str(row[0])[:20],fontName=font(),fontSize=9,fillColor=HexColor("#285F95")))
            text=str(row[1])[:92]
            n=max(12,int((boxw-16)/8))
            for line in range(4):
                drawing.add(String(x+8,y+57-line*11,text[line*n:(line+1)*n],fontName=font(),fontSize=8))
            drawing.add(String(x+8,y+6,str(row[2]),fontName=font(),fontSize=7,fillColor=HexColor("#586579")))
            if i%2==0 and i+1<len(rows):
                drawing.add(Line(x+boxw,y+44,x+boxw+15,y+44,strokeColor=HexColor("#285F95")))
        return drawing
    drawing = Drawing(width, height)
    values, labels = visual["values"], visual["labels"]
    lo, hi = min(0,min(values)), max(0,max(values))
    if hi == lo:
        hi = lo+1
    left, bottom, top = 70, 40, height-20
    plotw = width-left-15
    y = lambda val: bottom + (val-lo)/(hi-lo)*(top-bottom)
    drawing.add(Line(left,y(0),width-15,y(0),strokeColor=HexColor("#A5B2C3")))
    for tick in range(5):
        val = lo+(hi-lo)*tick/4
        drawing.add(String(left-7,y(val)-3,f"{val:,.1f}",textAnchor="end",fontName=font(),fontSize=8))
        drawing.add(Line(left,y(val),width-15,y(val),strokeColor=HexColor("#E1E7EF"),strokeWidth=.4))
    previous = None
    for i,(value,label) in enumerate(zip(values,labels)):
        x = left+plotw*(i+.5)/len(values)
        if visual["kind"] == "line":
            if previous:
                drawing.add(Line(previous[0],previous[1],x,y(value),strokeColor=HexColor("#285F95"),strokeWidth=1.8))
            drawing.add(Circle(x,y(value),2.8,fillColor=HexColor("#285F95"),strokeColor=None))
            previous=(x,y(value))
        else:
            bw=plotw/len(values)*.55
            drawing.add(Rect(x-bw/2,min(y(0),y(value)),bw,abs(y(value)-y(0)),fillColor=HexColor("#285F95"),strokeColor=None))
        drawing.add(String(x,y(value)+5,f"{value:,.1f}",textAnchor="middle",fontName=font(),fontSize=8))
        # Two-line category labels; values remain fully available in the evidence workbook.
        short = str(label)[-22:]
        drawing.add(String(x,23,short[:11],textAnchor="middle",fontName=font(),fontSize=7.5))
        drawing.add(String(x,12,short[11:],textAnchor="middle",fontName=font(),fontSize=7.5))
    drawing.add(String(0,height-8,str(visual.get("unit","")),fontName=font(),fontSize=8))
    return drawing


def visual_png(visual):
    from reportlab.graphics import renderPDF
    import pypdfium2 as pdfium
    raw=renderPDF.drawToString(chart_drawing(visual))
    pdf=pdfium.PdfDocument(raw)
    page=pdf[0]
    bitmap=page.render(scale=2)
    out=io.BytesIO()
    bitmap.to_pil().save(out,format="PNG")
    bitmap.close(); page.close(); pdf.close()
    return out.getvalue()


def pdf_bytes(package):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, KeepInFrame
    from reportlab.pdfgen.canvas import Canvas
    from pypdf import PdfReader
    kai = font()
    width,height = A4
    usable = width-92
    blue = colors.HexColor("#1C5080")
    body = ParagraphStyle("body",fontName=kai,fontSize=10.3,leading=16.5,spaceAfter=9,wordWrap="CJK")
    heading = ParagraphStyle("heading",parent=body,fontSize=17,leading=23,textColor=blue,spaceAfter=14)
    small = ParagraphStyle("source",parent=body,fontSize=7.3,leading=10.5,textColor=colors.HexColor("#5D6673"),spaceAfter=9)
    takeaway = ParagraphStyle("takeaway",parent=body,fontSize=11,leading=17,textColor=blue,spaceAfter=14)
    def p(value,style=body):
        return Paragraph(escape(str(value)).replace("\n","<br/>"),style)
    def table(headers,rows):
        count=len(headers)
        widths = ([usable*.25,usable*.16,usable*.59] if count==3 and headers[1]=="取值" else
                  [usable*.2,usable*.65,usable*.15] if count==3 else
                  [usable*.31]+[usable*.69/(count-1)]*(count-1))
        cell=ParagraphStyle("cell",parent=body,fontSize=8,leading=11,spaceAfter=0)
        head=ParagraphStyle("th",parent=cell,textColor=colors.white)
        data=[[p(v,head) for v in headers]]+[[p(row[i] if i<len(row) else "",cell) for i in range(count)] for row in rows]
        t=Table(data,colWidths=widths,repeatRows=1,hAlign="LEFT")
        t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),blue), ("VALIGN",(0,0),(-1,-1),"TOP"),
            ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F0F4F8")]),
            ("LINEBELOW",(0,0),(-1,0),.6,blue), ("LINEBELOW",(0,-1),(-1,-1),.6,blue),
            ("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),
            ("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5)]))
        return t
    story=[]
    story += [Spacer(1,35),p(package["target"],ParagraphStyle("cover",parent=heading,fontSize=30,leading=40)),
              p(package["kind"]+"研究 · "+package["report_type"],heading),p(package["status"],takeaway),
              p("目标观察期："+package["period"]),p("生成时间："+package["created_at"][:19]),Spacer(1,18)]
    if package["chapters"]:
        story += [p("核心判断",heading),p(package["chapters"][0]["takeaway"]),p(package["chapters"][0]["paragraphs"][0])]
    story += [p("阅读口径",heading),p("正文引用 [S编号] 对应文末来源与证据工作簿。图表下方标注来源。披露事实、研究推断与预测假设分别说明；检索时间不等于财报期。公开资料不能证明覆盖全部最新信息。"),
              p("财务模型："+package["model"]["unit"]+"；"+package["model"]["opening_note"]),PageBreak(),p("目录与图表索引",heading)]
    for i,c in enumerate(package["chapters"],1):
        story += [p(f"{i:02d}  {c['title']}",ParagraphStyle("toc",parent=body,fontSize=9.5,leading=12,spaceAfter=3))]
    story += [Spacer(1,12),p("每章对应图表编号与章节编号一致；实际页码连续显示在页脚。",small),PageBreak()]
    # Explicit chapter breaks ensure >=20 substantive content pages, excluding front/back matter.
    for i,c in enumerate(package["chapters"],1):
        visual=c["visual"]
        chapter_story = [p(f"{i:02d}  {c['title']}",heading),p(c["takeaway"],takeaway)]
        paragraphs=c["paragraphs"]
        split=max(2,len(paragraphs)//2)
        chapter_story += [p(t) for t in paragraphs[:split]]
        graph = table(visual["headers"],visual["rows"]) if visual["kind"]=="table" else chart_drawing(visual,usable,185)
        # Chart and source cannot be orphaned; narrative before and after relates to this exact exhibit.
        chapter_story += [p(f"图表 {i}  {visual['title']}",takeaway), graph,Spacer(1,5),p(visual["source"],small)]
        chapter_story += [p(t) for t in paragraphs[split:]]
        # Fit a complete argument and exhibit to one page; no almost-empty overflow page.
        story += [KeepInFrame(usable,height-114,chapter_story,mode="shrink",hAlign="LEFT",vAlign="TOP")]
        story += [PageBreak()]
    story += [p("来源与取数审计",heading),p("以下为实际打开的资料；发表日期与财务期间以原文为准，检索未覆盖不等于未发生。",small)]
    for d in package["sources"]:
        story += [p(f"[{d['id']}] {d['title']}",takeaway),p(d["url"],small),
                  p(f"方法：{d['method']}；抓取时间：{d['retrieved_at']}",small)]
    story += [p("质量检查与模型边界",heading)]
    for label,passed in package["checks"].items():
        story.append(p(("通过：" if passed else "未通过：")+label))
    story += [p(package["model"]["note"]),p("本报告由自动化研究工具生成。结构检查不等同于事实或投资结论认证；请复核引文语境、会计口径、预测假设与资料覆盖。",small)]
    class NumberedCanvas(Canvas):
        def __init__(self,*a,**kw):
            super().__init__(*a,**kw)
            self.saved=[]
        def showPage(self):
            self.saved.append(dict(self.__dict__))
            self._startPage()
        def save(self):
            total=len(self.saved)
            for state in self.saved:
                self.__dict__.update(state)
                self.setFont(kai,8)
                self.setFillColor(blue)
                self.drawString(46,height-28,str(package["target"])[:38]+" · "+package["report_type"])
                self.setStrokeColor(blue)
                self.line(46,height-36,width-46,height-36)
                self.setFillColor(colors.HexColor("#586579"))
                self.drawString(46,26,"自动研究 · 披露数据与假设分列 · 详见来源和证据")
                self.drawRightString(width-46,26,f"{self._pageNumber} / {total}")
                Canvas.showPage(self)
            Canvas.save(self)
    output=io.BytesIO()
    SimpleDocTemplate(output,pagesize=A4,rightMargin=46,leftMargin=46,topMargin=55,bottomMargin=47,
                      title=package["target"]+"研究报告",author="Research Agent").build(story,canvasmaker=NumberedCanvas)
    raw=output.getvalue()
    reader=PdfReader(io.BytesIO(raw))
    min_content=6 if package["report_type"] in ("周报","季报") else 20
    if len(package["chapters"]) < min_content or len(reader.pages) < min_content+3:
        raise ValueError("实际 PDF 未达到最低内容页数，未提供正式报告下载")
    return raw


def pptx_bytes(package,pdf):
    """Render each PDF page onto a portrait slide: fonts remain Kai on every computer."""
    import pypdfium2 as pdfium
    from pptx import Presentation
    from pptx.util import Inches
    prs=Presentation()
    prs.slide_width=Inches(8.27)
    prs.slide_height=Inches(11.69)
    document=pdfium.PdfDocument(pdf)
    try:
        for i in range(len(document)):
            page=document[i]
            bitmap=page.render(scale=1.65)
            try:
                im=bitmap.to_pil()
                buf=io.BytesIO()
                im.save(buf,format="PNG")
                slide=prs.slides.add_slide(prs.slide_layouts[6])
                slide.shapes.add_picture(io.BytesIO(buf.getvalue()),0,0,width=prs.slide_width,height=prs.slide_height)
                slide.notes_slide.notes_text_frame.text="楷体保真版；可编辑正文见同次 Word 文件。"+package["target"]
            finally:
                bitmap.close()
                page.close()
    finally:
        document.close()
    output=io.BytesIO()
    prs.save(output)
    return output.getvalue()


def apply_docx_font(document):
    from docx.oxml.ns import qn
    for style in document.styles:
        if hasattr(style,"font"):
            style.font.name=FONT_FAMILY
            fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
            for attr in ("ascii","hAnsi","eastAsia","cs"):
                fonts.set(qn("w:"+attr),FONT_FAMILY)
    for fonts in document.element.xpath(".//w:rFonts"):
        for attr in ("ascii","hAnsi","eastAsia","cs"):
            fonts.set(qn("w:"+attr),FONT_FAMILY)
    for section in document.sections:
        for part in (section.header,section.footer):
            for paragraph in part.paragraphs:
                for run in paragraph.runs:
                    run.font.name=FONT_FAMILY


def docx_bytes(package):
    from docx import Document
    from docx.shared import Pt,Inches
    from reportlab.graphics import renderPDF
    import pypdfium2 as pdfium
    doc=Document()
    doc.styles["Normal"].font.size=Pt(10.5)
    doc.add_heading(package["target"]+" · "+package["report_type"],0)
    doc.add_paragraph(package["status"])
    doc.add_paragraph("预测假设与来源见附录；"+package["model"]["opening_note"])
    for i,c in enumerate(package["chapters"],1):
        doc.add_page_break()
        doc.add_heading(f"{i:02d} {c['title']}",1)
        doc.add_paragraph(c["takeaway"])
        half=max(2,len(c["paragraphs"])//2)
        for paragraph in c["paragraphs"][:half]:
            doc.add_paragraph(paragraph)
        v=c["visual"]
        doc.add_paragraph(f"图表{i}：{v['title']}")
        if v["kind"]=="table":
            tab=doc.add_table(rows=1,cols=len(v["headers"]))
            tab.style="Light Shading Accent 1"
            for cell,txt in zip(tab.rows[0].cells,v["headers"]): cell.text=str(txt)
            for row in v["rows"]:
                for cell,txt in zip(tab.add_row().cells,row): cell.text=str(txt)
        else:
            data=renderPDF.drawToString(chart_drawing(v))
            pdf=pdfium.PdfDocument(data)
            page=pdf[0]
            bm=page.render(scale=2)
            buf=io.BytesIO()
            bm.to_pil().save(buf,format="PNG")
            doc.add_picture(io.BytesIO(buf.getvalue()),width=Inches(6.2))
            bm.close(); page.close(); pdf.close()
        doc.add_paragraph(v["source"],style="Caption")
        for paragraph in c["paragraphs"][half:]: doc.add_paragraph(paragraph)
    doc.add_page_break()
    doc.add_heading("来源索引",1)
    for s in package["sources"]: doc.add_paragraph(f"[{s['id']}] {s['title']}\n{s['url']}\n{s['retrieved_at']}")
    apply_docx_font(doc)
    out=io.BytesIO(); doc.save(out)
    return out.getvalue()


def evidence_bytes(package):
    """Auditable workbook including source quotes and editable formula-driven Pro Forma."""
    import xlsxwriter
    out=io.BytesIO()
    wb=xlsxwriter.Workbook(out,{"in_memory":True,"strings_to_formulas":False,"strings_to_urls":False})
    fmt=wb.add_format({"font_name":FONT_FAMILY,"font_size":10,"num_format":"#,##0.00"})
    header=wb.add_format({"font_name":FONT_FAMILY,"bold":True,"bg_color":"#1C5080","font_color":"#FFFFFF"})
    wrap=wb.add_format({"font_name":FONT_FAMILY,"text_wrap":True,"valign":"top"})
    def sheet(name,headers,rows):
        ws=wb.add_worksheet(name)
        ws.write_row(0,0,headers,header)
        for i,row in enumerate(rows,1): ws.write_row(i,0,row,fmt)
        ws.freeze_panes(1,1); ws.set_column(0,len(headers)-1,22,fmt)
        return ws
    model=package["model"]
    keys=list(model["assumptions"])
    inputs=sheet("Assumptions",["驱动项","值（可编辑）","依据"],[[k,model["assumptions"][k],model["assumption_basis"][k]] for k in keys])
    inputs.write_row(len(keys)+2,0,["基期收入",model["base_revenue"],model["base_source"]],fmt)
    inputs.write_row(len(keys)+3,0,["计量单位",model["unit"],model["opening_note"]],wrap)
    mapping={k:f"Assumptions!$B${i+2}" for i,k in enumerate(keys)}
    rows=[("营业收入","revenue"),("成本","cost"),("毛利润","gross_profit"),("经营费用","opex"),("EBIT","ebit"),
          ("折旧摊销","da"),("利息","interest"),("所得税","tax"),("净利润","net_income"),("现金","cash"),
          ("应收账款","receivables"),("存货","inventory"),("固定资产净额","ppe"),("资产总额","assets"),
          ("应付账款","payables"),("有息负债","debt"),("负债总额","liabilities"),("权益","equity"),
          ("营运资金增加","nwc_delta"),("经营现金流","cfo"),("资本开支","capex"),("分红","dividends"),
          ("新增借款","funding"),("投资现金流","cfi"),("筹资现金流","cff"),("现金净增加","cash_change"),
          ("FCFF","fcff"),("资产负债平衡检查","balance_check"),("现金流检查","cash_check")]
    ws=sheet("Pro Forma",["项目", "基期/期初结构假设"]+[r["year"] for r in model["rows"]],[])
    rowmap={k:i+2 for i,(_,k) in enumerate(rows)}
    for i,(label,key) in enumerate(rows,1):
        ws.write(i,0,label,fmt)
        ws.write(i,1,model["base_revenue"] if key=="revenue" else model["opening"].get(key,0),fmt)
    ws.write_formula(rowmap["revenue"]-1,1,f"=Assumptions!B{len(keys)+3}",fmt,model["base_revenue"])
    # Opening cells are editable and labelled assumptions; equity and assets must stay balanced after edits.
    ws.write_formula(rowmap["assets"]-1,1,"=SUM(B11:B14)",fmt,sum(model["opening"][k] for k in ("cash","receivables","inventory","ppe")))
    ws.write_formula(rowmap["liabilities"]-1,1,"=B16+B17",fmt,model["opening"]["payables"]+model["opening"]["debt"])
    ws.write_formula(rowmap["equity"]-1,1,"=B15-B18",fmt,model["opening"]["equity"])
    from xlsxwriter.utility import xl_col_to_name
    for i,result in enumerate(model["rows"],2):
        col,prev=xl_col_to_name(i),xl_col_to_name(i-1)
        r=lambda k: col+str(rowmap[k])
        p=lambda k: prev+str(rowmap[k])
        a=lambda k: mapping[k]
        f={
         "revenue":f"{p('revenue')}*(1+{a('growth')})", "cost":f"{r('revenue')}*(1-{a('gross_margin')})",
         "gross_profit":f"{r('revenue')}-{r('cost')}","opex":f"{r('revenue')}*{a('opex_ratio')}",
         "ebit":f"{r('gross_profit')}-{r('opex')}","da":f"MIN({r('revenue')}*{a('da_ratio')},{p('ppe')}+{r('capex')})",
         "interest":f"{p('debt')}*{a('interest_rate')}","tax":f"MAX({r('ebit')}-{r('interest')},0)*{a('tax_rate')}",
         "net_income":f"{r('ebit')}-{r('interest')}-{r('tax')}","receivables":f"{r('revenue')}*{a('ar_ratio')}",
         "inventory":f"{r('revenue')}*{a('inventory_ratio')}","payables":f"{r('revenue')}*{a('ap_ratio')}",
         "nwc_delta":f"{r('receivables')}+{r('inventory')}-{r('payables')}-{p('receivables')}-{p('inventory')}+{p('payables')}",
         "cfo":f"{r('net_income')}+{r('da')}-{r('nwc_delta')}","capex":f"{r('revenue')}*{a('capex_ratio')}",
         "dividends":f"MAX({r('net_income')},0)*{a('payout')}","funding":f"MAX(0,-({p('cash')}+{r('cfo')}-{r('capex')}-{r('dividends')}))",
         "cash":f"{p('cash')}+{r('cash_change')}","debt":f"{p('debt')}+{r('funding')}","ppe":f"{p('ppe')}+{r('capex')}-{r('da')}",
         "equity":f"{p('equity')}+{r('net_income')}-{r('dividends')}","assets":f"{r('cash')}+{r('receivables')}+{r('inventory')}+{r('ppe')}",
         "liabilities":f"{r('payables')}+{r('debt')}","cfi":f"-{r('capex')}","cff":f"{r('funding')}-{r('dividends')}",
         "cash_change":f"{r('cfo')}+{r('cfi')}+{r('cff')}","fcff":f"{r('ebit')}*(1-{a('tax_rate')})+{r('da')}-{r('capex')}-{r('nwc_delta')}",
         "balance_check":f"{r('assets')}-{r('liabilities')}-{r('equity')}","cash_check":f"{r('cash')}-{p('cash')}-{r('cash_change')}"}
        for key,formula in f.items(): ws.write_formula(rowmap[key]-1,i,"="+formula,fmt,result[key])
    equity=sheet("Equity",["权益变动"]+[r["year"] for r in model["rows"]],[])
    labels=[("期初权益","opening_equity"),("净利润","net_income"),("增发（无交易假设）","equity_issuance"),
            ("回购（无交易假设）","equity_buybacks"),("分红","dividends"),("期末权益","equity"),("权益勾稽","equity_check")]
    for i,(label,key) in enumerate(labels,1):
        equity.write(i,0,label,fmt)
        for col,result in enumerate(model["rows"],1):
            c=xl_col_to_name(col); pc=xl_col_to_name(col+1); prev=xl_col_to_name(col)
            formulas={"opening_equity":f"'Pro Forma'!{prev}19", "net_income":f"'Pro Forma'!{pc}10",
                      "equity_issuance":"0", "equity_buybacks":"0", "dividends":f"'Pro Forma'!{pc}23",
                      "equity":f"{c}2+{c}3+{c}4-{c}5-{c}6", "equity_check":f"{c}7-'Pro Forma'!{pc}19"}
            equity.write_formula(i,col,"="+formulas[key],fmt,result[key])
    sensitivity=sheet("DCF Sensitivity",["WACC / g",.015,.02,.025,.03,.035],[])
    for row,w in enumerate([.08,.09,.10,.11,.12],1):
        sensitivity.write(row,0,w,fmt)
        for col,g in enumerate([.015,.02,.025,.03,.035],1):
            c=xl_col_to_name(col)
            formula=f"=NPV($A{row+1},'Pro Forma'!C28:G28)+'Pro Forma'!G28*(1+{c}$1)/($A{row+1}-{c}$1)/(1+$A{row+1})^5"
            sensitivity.write_formula(row,col,formula,fmt,model["sensitivity"][row-1][col-1])
    sensitivity.conditional_format("B2:F6",{"type":"3_color_scale"})
    sheet("Evidence",["主体","事实","期间","指标","原文数值","单位","来源","逐字引文"],
          [[f["subject"],f["claim"],f["period"],f["metric"],f.get("value",""),f["unit"],f["source_id"],f["quote"]] for f in package["facts"]])
    sheet("Sources",["编号","标题","URL","抓取时间","方法"],[[d["id"],d["title"],d["url"],d["retrieved_at"],d["method"]] for d in package["sources"]])
    sheet("Checks",["检查","结果"],[[k,"通过" if v else "未通过"] for k,v in package["checks"].items()])
    for style in wb.formats: style.set_font_name(FONT_FAMILY)
    wb.close()
    return out.getvalue()
